import logging
from docx import Document
import re
import os
import uuid
from typing import List, Dict, Any

logger = logging.getLogger(__name__)

class DocxParser:
    def __init__(self, file_stream, upload_dir: str = "static/uploads"):
        self.doc = Document(file_stream)
        self.upload_dir = upload_dir

    def _extract_images_from_paragraph(self, para) -> List[str]:
        para_images = []
        try:
            xml_str = para._p.xml
            embeds = re.findall(r'r:embed="([^"]+)"', xml_str)
            for rId in embeds:
                if rId in para.part.related_parts:
                    image_part = para.part.related_parts[rId]
                    image_bytes = image_part.blob
                    filename = f"{uuid.uuid4()}.png"
                    filepath = os.path.join(self.upload_dir, filename)
                    with open(filepath, "wb") as f:
                        f.write(image_bytes)
                    para_images.append(f"/api/uploads/{filename}")
        except Exception as e:
            logger.error(f"Error extracting image from paragraph: {e}")
        return para_images

    def parse(self) -> List[Dict[str, Any]]:
        questions = []
        current_question = None

        os.makedirs(self.upload_dir, exist_ok=True)

        # Collect all items (paragraphs + table cells)
        elements = []
        for para in self.doc.paragraphs:
            elements.append((para.text.strip(), self._extract_images_from_paragraph(para)))

        # Also collect paragraphs inside tables if present
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        txt = para.text.strip()
                        if txt:
                            elements.append((txt, self._extract_images_from_paragraph(para)))

        # Regex patterns for flexible matching
        question_pattern = re.compile(r'^(câu|câu hỏi|bài|question)?s*d+[.:-]', re.IGNORECASE)
        option_pattern = re.compile(r'^[A-D][.:)]', re.IGNORECASE)

        for text, para_images in elements:
            if not text and not para_images:
                continue

            if question_pattern.match(text):
                if current_question:
                    questions.append(current_question)

                clean_content = question_pattern.sub('', text).strip()
                current_question = {
                    "content": clean_content,
                    "options": [],
                    "correct_option": 0,
                    "subject": "General",
                    "grade_level": 10,
                    "difficulty": "THONG_HIEU",
                    "explanation": "",
                    "image_url": para_images[0] if para_images else None,
                    "media": {"images": para_images} if para_images else None
                }
            elif current_question and para_images:
                if not current_question.get("image_url"):
                    current_question["image_url"] = para_images[0]
                images_list = current_question.get("media", {}).get("images", []) if current_question.get("media") else []
                images_list.extend(para_images)
                if not current_question.get("media"):
                    current_question["media"] = {}
                current_question["media"]["images"] = images_list

            if current_question and text:
                if option_pattern.match(text):
                    option_text = option_pattern.sub('', text).strip()
                    current_question["options"].append(option_text)
                elif text.lower().startswith("answer:") or text.lower().startswith("đáp án:"):
                    ans_str = text.split(":")[-1].strip().upper()
                    mapping = {'A': 0, 'B': 1, 'C': 2, 'D': 3}
                    current_question["correct_option"] = mapping.get(ans_str, 0)
                elif text.lower().startswith("explanation:") or text.lower().startswith("lời giải:"):
                    current_question["explanation"] = text.split(":", 1)[-1].strip()
                elif not question_pattern.match(text) and not option_pattern.match(text):
                    if text and not text.lower().startswith("answer:") and not text.lower().startswith("đáp án:") and not text.lower().startswith("explanation:") and not text.lower().startswith("lời giải:"):
                        if current_question["content"]:
                            current_question["content"] += f"\n{text}"

        if current_question:
            questions.append(current_question)

        return questions
